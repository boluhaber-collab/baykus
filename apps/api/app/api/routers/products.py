"""Products + variants + stock movements CRUD."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from sqlalchemy.orm import Session, joinedload

from app.core.deps import get_db, require_roles
from app.models.product import DEFAULT_WAREHOUSE, Product, ProductVariant, StockMovement
from app.models.price_list import PriceList, PriceListItem
from app.models.user import User
from app.schemas.product import (
    CriticalStockItem,
    ProductCreate,
    ProductDetail,
    ProductListItem,
    ProductPriceListRef,
    ProductPricingOut,
    ProductUpdate,
    StockAdjustIn,
    StockAdjustOut,
    StockMovementOut,
    VariantCreate,
    VariantOut,
    VariantUpdate,
)

router = APIRouter(prefix="/products", tags=["products"])

READ_ROLES = ("admin", "satış", "üretim", "muhasebe")
WRITE_ROLES = ("admin", "üretim", "satış")
STOCK_ROLES = ("admin", "üretim")


def _variant_stock_sum(product: Product) -> int:
    if product.variants:
        return sum(v.stock_qty for v in product.variants)
    return product.stock_qty or 0


def _is_critical(product: Product) -> bool:
    if product.product_type == "hizmet":
        return False
    threshold = product.critical_stock_threshold or 0
    if product.variants:
        return any(v.stock_qty < threshold for v in product.variants)
    return (product.stock_qty or 0) < threshold


def _sync_product_stock(product: Product) -> None:
    """Keep product.stock_qty aligned with variant sum when variants exist."""
    if product.variants:
        product.stock_qty = sum(v.stock_qty for v in product.variants)


def _variant_out(v: ProductVariant, threshold: int) -> VariantOut:
    return VariantOut(
        id=v.id,
        product_id=v.product_id,
        name=v.name,
        sku=v.sku,
        color=v.color,
        size=v.size,
        print_type=v.print_type,
        barcode=v.barcode,
        price=v.price,
        stock_qty=v.stock_qty,
        is_critical=v.stock_qty < threshold,
    )


def _movement_out(m: StockMovement) -> StockMovementOut:
    return StockMovementOut(
        id=m.id,
        product_id=m.product_id,
        variant_id=m.variant_id,
        direction=m.direction,
        quantity=m.quantity,
        qty_before=m.qty_before,
        qty_after=m.qty_after,
        reason=m.reason,
        note=m.note,
        warehouse=m.warehouse,
        created_by_user_id=m.created_by_user_id,
        created_at=m.created_at,
        variant_sku=m.variant.sku if m.variant else None,
        variant_name=m.variant.name if m.variant else None,
    )


def _list_item(product: Product) -> ProductListItem:
    total = _variant_stock_sum(product)
    return ProductListItem(
        id=product.id,
        sku=product.sku,
        name=product.name,
        category=product.category,
        brand=product.brand,
        product_type=product.product_type,
        base_price=product.base_price,
        purchase_price=product.purchase_price or Decimal("0"),
        cost=product.cost or Decimal("0"),
        stock_qty=product.stock_qty or 0,
        total_stock=total,
        critical_stock_threshold=product.critical_stock_threshold or 0,
        is_critical=_is_critical(product),
        is_active=bool(product.is_active),
        warehouse=product.warehouse,
        variants_count=len(product.variants or []),
        photo_url=product.photo_url,
    )


def _detail(product: Product, movements: list[StockMovement] | None = None) -> ProductDetail:
    threshold = product.critical_stock_threshold or 0
    return ProductDetail(
        id=product.id,
        sku=product.sku,
        name=product.name,
        category=product.category,
        brand=product.brand,
        supplier_name=product.supplier_name,
        product_type=product.product_type,
        description=product.description,
        base_price=product.base_price,
        purchase_price=product.purchase_price or Decimal("0"),
        cost=product.cost or Decimal("0"),
        photo_url=product.photo_url,
        is_active=bool(product.is_active),
        critical_stock_threshold=threshold,
        warehouse=product.warehouse,
        stock_qty=product.stock_qty or 0,
        total_stock=_variant_stock_sum(product),
        is_critical=_is_critical(product),
        created_at=product.created_at,
        updated_at=product.updated_at,
        variants=[_variant_out(v, threshold) for v in (product.variants or [])],
        recent_movements=[_movement_out(m) for m in (movements or [])],
    )


def _get_product(db: Session, product_id: int) -> Product:
    product = (
        db.query(Product)
        .options(joinedload(Product.variants))
        .filter(Product.id == product_id)
        .first()
    )
    if not product:
        raise HTTPException(status_code=404, detail="Ürün bulunamadı")
    return product


def _ensure_unique_sku(db: Session, sku: str, exclude_product_id: int | None = None) -> None:
    q = db.query(Product).filter(Product.sku == sku)
    if exclude_product_id is not None:
        q = q.filter(Product.id != exclude_product_id)
    if q.first():
        raise HTTPException(status_code=400, detail=f"SKU zaten kullanılıyor: {sku}")


def _ensure_unique_variant_sku(
    db: Session, sku: str, exclude_variant_id: int | None = None
) -> None:
    q = db.query(ProductVariant).filter(ProductVariant.sku == sku)
    if exclude_variant_id is not None:
        q = q.filter(ProductVariant.id != exclude_variant_id)
    if q.first():
        raise HTTPException(status_code=400, detail=f"Varyant SKU zaten kullanılıyor: {sku}")


@router.get("", response_model=list[ProductListItem])
def list_products(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
    q: str | None = Query(default=None),
    category: str | None = Query(default=None),
    product_type: str | None = Query(default=None, alias="type"),
    critical_only: bool = Query(default=False),
    active_only: bool = Query(default=False),
    skip: int = 0,
    limit: int = 200,
) -> list[ProductListItem]:
    query = db.query(Product).options(joinedload(Product.variants))
    if q:
        like = f"%{q}%"
        query = query.filter(
            (Product.name.ilike(like))
            | (Product.sku.ilike(like))
            | (Product.brand.ilike(like))
            | (Product.category.ilike(like))
        )
    if category:
        query = query.filter(Product.category == category)
    if product_type:
        query = query.filter(Product.product_type == product_type)
    if active_only:
        query = query.filter(Product.is_active.is_(True))
    products = query.order_by(Product.id.desc()).offset(skip).limit(limit).all()
    items = [_list_item(p) for p in products]
    if critical_only:
        items = [i for i in items if i.is_critical]
    return items


@router.get("/categories", response_model=list[str])
def list_categories(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
) -> list[str]:
    rows = (
        db.query(Product.category)
        .filter(Product.category.isnot(None), Product.category != "")
        .distinct()
        .order_by(Product.category)
        .all()
    )
    return [r[0] for r in rows if r[0]]


@router.get("/stock-summary")
def stock_summary(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*STOCK_ROLES, "satış")),
) -> dict:
    products = db.query(Product).options(joinedload(Product.variants)).all()
    critical: list[dict] = []
    for p in products:
        if p.product_type == "hizmet" or not p.is_active:
            continue
        threshold = p.critical_stock_threshold or 0
        if p.variants:
            for v in p.variants:
                if v.stock_qty < threshold:
                    critical.append(
                        {
                            "product_id": p.id,
                            "sku": v.sku,
                            "name": f"{p.name} / {v.name}",
                            "qty": v.stock_qty,
                            "threshold": threshold,
                        }
                    )
        elif (p.stock_qty or 0) < threshold:
            critical.append(
                {
                    "product_id": p.id,
                    "sku": p.sku,
                    "name": p.name,
                    "qty": p.stock_qty,
                    "threshold": threshold,
                }
            )
    return {
        "total_skus": len(products),
        "active_skus": sum(1 for p in products if p.is_active),
        "critical_count": len(critical),
        "low_stock": critical,
    }


@router.get("/critical", response_model=list[CriticalStockItem])
def list_critical_stock(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*STOCK_ROLES, "satış")),
) -> list[CriticalStockItem]:
    products = (
        db.query(Product)
        .options(joinedload(Product.variants))
        .filter(Product.is_active.is_(True), Product.product_type != "hizmet")
        .order_by(Product.name)
        .all()
    )
    items: list[CriticalStockItem] = []
    for p in products:
        threshold = p.critical_stock_threshold or 0
        if p.variants:
            for v in p.variants:
                if v.stock_qty < threshold:
                    items.append(
                        CriticalStockItem(
                            product_id=p.id,
                            product_sku=p.sku,
                            product_name=p.name,
                            category=p.category,
                            variant_id=v.id,
                            variant_sku=v.sku,
                            variant_name=v.name,
                            stock_qty=v.stock_qty,
                            critical_stock_threshold=threshold,
                            warehouse=p.warehouse,
                        )
                    )
        elif (p.stock_qty or 0) < threshold:
            items.append(
                CriticalStockItem(
                    product_id=p.id,
                    product_sku=p.sku,
                    product_name=p.name,
                    category=p.category,
                    variant_id=None,
                    variant_sku=None,
                    variant_name=None,
                    stock_qty=p.stock_qty or 0,
                    critical_stock_threshold=threshold,
                    warehouse=p.warehouse,
                )
            )
    return items




@router.get("/stock/export")
def export_stock_csv(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
    fmt: str = Query(default="csv", pattern="^(csv|xlsx)$"),
):
    """Export products/variants stock as CSV or XLSX template."""
    import csv
    import io

    products = (
        db.query(Product)
        .options(joinedload(Product.variants))
        .order_by(Product.name)
        .all()
    )
    headers = [
        "sku", "name", "category", "product_type", "warehouse",
        "stock_qty", "critical_stock_threshold", "base_price", "purchase_price",
        "variant_sku", "variant_name", "color", "size", "variant_stock",
    ]
    rows: list[list] = []
    for prod in products:
        if prod.variants:
            for v in prod.variants:
                rows.append([
                    prod.sku, prod.name, prod.category or "", prod.product_type,
                    prod.warehouse or DEFAULT_WAREHOUSE,
                    prod.stock_qty or 0, prod.critical_stock_threshold or 10,
                    float(prod.base_price or 0), float(prod.purchase_price or 0),
                    v.sku, v.name, v.color or "", v.size or "", v.stock_qty or 0,
                ])
        else:
            rows.append([
                prod.sku, prod.name, prod.category or "", prod.product_type,
                prod.warehouse or DEFAULT_WAREHOUSE,
                prod.stock_qty or 0, prod.critical_stock_threshold or 10,
                float(prod.base_price or 0), float(prod.purchase_price or 0),
                "", "", "", "", "",
            ])

    if fmt == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="openpyxl yüklü değil") from exc
        wb = Workbook()
        ws = wb.active
        ws.title = "Stok"
        ws.append(headers)
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="stok-export.xlsx"'},
        )

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return Response(
        content=buf.getvalue().encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="stok-export.csv"'},
    )


@router.get("/stock/import-template")
def stock_import_template(
    _: User = Depends(require_roles(*READ_ROLES)),
    fmt: str = Query(default="csv", pattern="^(csv|xlsx)$"),
):
    """Empty template for stock import."""
    import csv
    import io

    headers = [
        "sku", "name", "category", "product_type", "warehouse",
        "stock_qty", "critical_stock_threshold", "base_price", "purchase_price",
    ]
    sample = ["ORNK-001", "Örnek Ürün", "Tekstil", "stoklu", "Ana Depo", "10", "5", "100", "60"]
    if fmt == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="openpyxl yüklü değil") from exc
        wb = Workbook()
        ws = wb.active
        ws.title = "Stok"
        ws.append(headers)
        ws.append(sample)
        buf = io.BytesIO()
        wb.save(buf)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="stok-sablon.xlsx"'},
        )
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerow(sample)
    return Response(
        content=buf.getvalue().encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="stok-sablon.csv"'},
    )


@router.post("/stock/import")
async def import_stock_csv(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*STOCK_ROLES)),
) -> dict:
    """Import/update products from CSV (sku key). Creates missing SKUs; updates stock/prices."""
    import csv
    import io

    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="Boş dosya")
    name = (file.filename or "").lower()
    rows: list[dict] = []
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="openpyxl yüklü değil") from exc
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip() for h in next(it)]
        for vals in it:
            if not vals or all(v is None or str(v).strip() == "" for v in vals):
                continue
            rows.append({headers[i]: vals[i] for i in range(min(len(headers), len(vals)))})
    else:
        text_data = raw.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text_data))
        rows = list(reader)

    created = updated = skipped = 0
    errors: list[str] = []
    for i, row in enumerate(rows, start=2):
        sku = str(row.get("sku") or "").strip()
        if not sku:
            skipped += 1
            continue
        try:
            prod = db.query(Product).filter(Product.sku == sku).first()
            stock_qty = int(float(row.get("stock_qty") or 0))
            thr = int(float(row.get("critical_stock_threshold") or 10))
            base = Decimal(str(row.get("base_price") or 0))
            purchase = Decimal(str(row.get("purchase_price") or 0))
            warehouse = str(row.get("warehouse") or DEFAULT_WAREHOUSE).strip() or DEFAULT_WAREHOUSE
            pname = str(row.get("name") or sku).strip()
            category = str(row.get("category") or "").strip() or None
            ptype = str(row.get("product_type") or "stoklu").strip() or "stoklu"
            if prod:
                before = int(prod.stock_qty or 0)
                prod.name = pname
                prod.category = category
                prod.product_type = ptype
                prod.warehouse = warehouse
                prod.critical_stock_threshold = thr
                prod.base_price = base
                prod.purchase_price = purchase
                if not prod.variants:
                    prod.stock_qty = stock_qty
                    if stock_qty != before:
                        direction = "increase" if stock_qty > before else "decrease"
                        db.add(
                            StockMovement(
                                product_id=prod.id,
                                direction=direction,
                                quantity=abs(stock_qty - before),
                                qty_before=before,
                                qty_after=stock_qty,
                                reason="import",
                                note="CSV/XLSX stok içe aktarma",
                                warehouse=warehouse,
                                created_by_user_id=user.id,
                            )
                        )
                updated += 1
            else:
                prod = Product(
                    sku=sku,
                    name=pname,
                    category=category,
                    product_type=ptype,
                    warehouse=warehouse,
                    stock_qty=stock_qty,
                    critical_stock_threshold=thr,
                    base_price=base,
                    purchase_price=purchase,
                    cost=purchase,
                    is_active=True,
                )
                db.add(prod)
                db.flush()
                if stock_qty:
                    db.add(
                        StockMovement(
                            product_id=prod.id,
                            direction="increase",
                            quantity=stock_qty,
                            qty_before=0,
                            qty_after=stock_qty,
                            reason="import",
                            note="CSV/XLSX stok içe aktarma (yeni)",
                            warehouse=warehouse,
                            created_by_user_id=user.id,
                        )
                    )
                created += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"Satır {i} ({sku}): {exc}")
    db.commit()
    return {
        "ok": True,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:20],
    }

@router.post("", response_model=ProductDetail, status_code=status.HTTP_201_CREATED)
def create_product(
    payload: ProductCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*WRITE_ROLES)),
) -> ProductDetail:
    if payload.product_type not in ("stoklu", "hizmet"):
        raise HTTPException(status_code=400, detail="product_type stoklu veya hizmet olmalı")
    _ensure_unique_sku(db, payload.sku)
    for v in payload.variants:
        _ensure_unique_variant_sku(db, v.sku)

    data = payload.model_dump(exclude={"variants"})
    product = Product(**data)
    db.add(product)
    db.flush()
    for v in payload.variants:
        db.add(ProductVariant(product_id=product.id, **v.model_dump()))
    db.flush()
    db.refresh(product)
    product = _get_product(db, product.id)
    _sync_product_stock(product)
    db.commit()
    product = _get_product(db, product.id)
    return _detail(product, [])


@router.get("/{product_id}", response_model=ProductDetail)
def get_product(
    product_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
) -> ProductDetail:
    product = _get_product(db, product_id)
    movements = (
        db.query(StockMovement)
        .options(joinedload(StockMovement.variant))
        .filter(StockMovement.product_id == product_id)
        .order_by(StockMovement.id.desc())
        .limit(50)
        .all()
    )
    return _detail(product, movements)


@router.put("/{product_id}", response_model=ProductDetail)
def update_product(
    product_id: int,
    payload: ProductUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*WRITE_ROLES)),
) -> ProductDetail:
    product = _get_product(db, product_id)
    data = payload.model_dump(exclude_unset=True, exclude={"variants"})
    if "sku" in data and data["sku"]:
        _ensure_unique_sku(db, data["sku"], exclude_product_id=product_id)
    if "product_type" in data and data["product_type"] not in (None, "stoklu", "hizmet"):
        raise HTTPException(status_code=400, detail="product_type stoklu veya hizmet olmalı")
    for key, value in data.items():
        setattr(product, key, value)
    product.updated_at = datetime.utcnow()

    if payload.variants is not None:
        # Replace variants: update matching sku, add new, remove missing
        existing_by_sku = {v.sku: v for v in product.variants}
        incoming_skus = set()
        for vdata in payload.variants:
            incoming_skus.add(vdata.sku)
            if vdata.sku in existing_by_sku:
                v = existing_by_sku[vdata.sku]
                for k, val in vdata.model_dump().items():
                    setattr(v, k, val)
            else:
                _ensure_unique_variant_sku(db, vdata.sku)
                db.add(ProductVariant(product_id=product.id, **vdata.model_dump()))
        for v in list(product.variants):
            if v.sku not in incoming_skus:
                db.delete(v)
        db.flush()
        product = _get_product(db, product_id)

    _sync_product_stock(product)
    db.commit()
    product = _get_product(db, product_id)
    movements = (
        db.query(StockMovement)
        .options(joinedload(StockMovement.variant))
        .filter(StockMovement.product_id == product_id)
        .order_by(StockMovement.id.desc())
        .limit(50)
        .all()
    )
    return _detail(product, movements)


@router.delete("/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_product(
    product_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin")),
) -> None:
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(status_code=404, detail="Ürün bulunamadı")
    db.delete(product)
    db.commit()


@router.post(
    "/{product_id}/variants",
    response_model=VariantOut,
    status_code=status.HTTP_201_CREATED,
)
def add_variant(
    product_id: int,
    payload: VariantCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*WRITE_ROLES)),
) -> VariantOut:
    product = _get_product(db, product_id)
    _ensure_unique_variant_sku(db, payload.sku)
    variant = ProductVariant(product_id=product.id, **payload.model_dump())
    db.add(variant)
    db.flush()
    _sync_product_stock(product)
    product.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(variant)
    return _variant_out(variant, product.critical_stock_threshold or 0)


@router.put("/{product_id}/variants/{variant_id}", response_model=VariantOut)
def update_variant(
    product_id: int,
    variant_id: int,
    payload: VariantUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*WRITE_ROLES)),
) -> VariantOut:
    product = _get_product(db, product_id)
    variant = db.get(ProductVariant, variant_id)
    if not variant or variant.product_id != product_id:
        raise HTTPException(status_code=404, detail="Varyant bulunamadı")
    data = payload.model_dump(exclude_unset=True)
    if "sku" in data and data["sku"]:
        _ensure_unique_variant_sku(db, data["sku"], exclude_variant_id=variant_id)
    for key, value in data.items():
        setattr(variant, key, value)
    _sync_product_stock(product)
    product.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(variant)
    return _variant_out(variant, product.critical_stock_threshold or 0)


@router.delete("/{product_id}/variants/{variant_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_variant(
    product_id: int,
    variant_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*WRITE_ROLES)),
) -> None:
    product = _get_product(db, product_id)
    variant = db.get(ProductVariant, variant_id)
    if not variant or variant.product_id != product_id:
        raise HTTPException(status_code=404, detail="Varyant bulunamadı")
    db.delete(variant)
    db.flush()
    product = _get_product(db, product_id)
    _sync_product_stock(product)
    product.updated_at = datetime.utcnow()
    db.commit()


@router.post("/{product_id}/stock/adjust", response_model=StockAdjustOut)
def adjust_stock(
    product_id: int,
    payload: StockAdjustIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*STOCK_ROLES)),
) -> StockAdjustOut:
    if payload.direction not in ("increase", "decrease"):
        raise HTTPException(status_code=400, detail="direction increase veya decrease olmalı")

    product = _get_product(db, product_id)
    if product.product_type == "hizmet":
        raise HTTPException(status_code=400, detail="Hizmet ürünlerinde stok hareketi yok")

    variant: ProductVariant | None = None
    if payload.variant_id is not None:
        variant = db.get(ProductVariant, payload.variant_id)
        if not variant or variant.product_id != product_id:
            raise HTTPException(status_code=404, detail="Varyant bulunamadı")
        qty_before = variant.stock_qty
    else:
        if product.variants:
            raise HTTPException(
                status_code=400,
                detail="Bu ürünün varyantları var; variant_id gerekli",
            )
        qty_before = product.stock_qty or 0

    if payload.direction == "increase":
        qty_after = qty_before + payload.quantity
    else:
        qty_after = qty_before - payload.quantity
        if qty_after < 0:
            raise HTTPException(
                status_code=400,
                detail=f"Yetersiz stok (mevcut: {qty_before})",
            )

    if variant is not None:
        variant.stock_qty = qty_after
    else:
        product.stock_qty = qty_after

    _sync_product_stock(product)
    product.updated_at = datetime.utcnow()

    movement = StockMovement(
        product_id=product.id,
        variant_id=variant.id if variant else None,
        direction=payload.direction,
        quantity=payload.quantity,
        qty_before=qty_before,
        qty_after=qty_after,
        reason=payload.reason,
        note=payload.note,
        warehouse=payload.warehouse or product.warehouse or DEFAULT_WAREHOUSE,
        created_by_user_id=user.id,
    )
    db.add(movement)
    db.commit()
    db.refresh(movement)
    product = _get_product(db, product_id)
    return StockAdjustOut(
        product_id=product.id,
        variant_id=variant.id if variant else None,
        direction=payload.direction,
        quantity=payload.quantity,
        qty_before=qty_before,
        qty_after=qty_after,
        total_stock=_variant_stock_sum(product),
        movement_id=movement.id,
    )


@router.get("/{product_id}/stock/movements", response_model=list[StockMovementOut])
def list_movements(
    product_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
    skip: int = 0,
    limit: int = 100,
) -> list[StockMovementOut]:
    if not db.get(Product, product_id):
        raise HTTPException(status_code=404, detail="Ürün bulunamadı")
    movements = (
        db.query(StockMovement)
        .options(joinedload(StockMovement.variant))
        .filter(StockMovement.product_id == product_id)
        .order_by(StockMovement.id.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return [_movement_out(m) for m in movements]


@router.get("/{product_id}/pricing", response_model=ProductPricingOut)
def product_pricing(
    product_id: int,
    variant_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
) -> ProductPricingOut:
    """Resolve unit price from active price list, else product/variant sale price; include stock."""
    product = (
        db.query(Product)
        .options(joinedload(Product.variants))
        .filter(Product.id == product_id)
        .first()
    )
    if not product:
        raise HTTPException(status_code=404, detail="Ürün bulunamadı")

    variant: ProductVariant | None = None
    if variant_id is not None:
        variant = next((v for v in product.variants if v.id == variant_id), None)
        if not variant:
            raise HTTPException(status_code=404, detail="Varyant bulunamadı")

    today = date.today()
    q = (
        db.query(PriceListItem, PriceList)
        .join(PriceList, PriceList.id == PriceListItem.price_list_id)
        .filter(PriceList.is_active.is_(True), PriceListItem.product_id == product_id)
    )
    if variant_id is not None:
        q = q.filter(
            (PriceListItem.variant_id == variant_id) | (PriceListItem.variant_id.is_(None))
        )
    else:
        q = q.filter(PriceListItem.variant_id.is_(None))
    rows = q.order_by(PriceList.id.desc(), PriceListItem.id.desc()).all()

    chosen_item = None
    chosen_list = None
    for item, pl in rows:
        if pl.valid_from and pl.valid_from > today:
            continue
        if pl.valid_to and pl.valid_to < today:
            continue
        if item.valid_from and item.valid_from > today:
            continue
        if item.valid_to and item.valid_to < today:
            continue
        chosen_item, chosen_list = item, pl
        break

    if chosen_item is not None and chosen_list is not None:
        unit_price = Decimal(chosen_item.unit_price or 0)
        source = "price_list"
        pl_id = chosen_list.id
        pl_name = chosen_list.name
    elif variant is not None:
        unit_price = Decimal(variant.price or 0)
        source = "variant"
        pl_id = None
        pl_name = None
    else:
        unit_price = Decimal(product.base_price or 0)
        source = "product"
        pl_id = None
        pl_name = None

    if variant is not None:
        stock_qty = int(variant.stock_qty or 0)
        name = f"{product.name} / {variant.name}"
        sku = variant.sku or product.sku
    else:
        stock_qty = _variant_stock_sum(product)
        name = product.name
        sku = product.sku

    thr = int(product.critical_stock_threshold or 10)
    return ProductPricingOut(
        product_id=product.id,
        variant_id=variant.id if variant else None,
        name=name,
        sku=sku,
        unit_price=unit_price,
        price_source=source,
        price_list_id=pl_id,
        price_list_name=pl_name,
        stock_qty=stock_qty,
        critical_stock_threshold=thr,
        is_critical=stock_qty < thr,
    )


@router.get("/{product_id}/price-lists", response_model=list[ProductPriceListRef])
def product_price_lists(
    product_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
) -> list[ProductPriceListRef]:
    if not db.get(Product, product_id):
        raise HTTPException(status_code=404, detail="Ürün bulunamadı")
    rows = (
        db.query(PriceListItem, PriceList)
        .join(PriceList, PriceList.id == PriceListItem.price_list_id)
        .filter(PriceListItem.product_id == product_id)
        .order_by(PriceList.name.asc())
        .all()
    )
    out: list[ProductPriceListRef] = []
    seen: set[int] = set()
    for item, pl in rows:
        if pl.id in seen:
            continue
        seen.add(pl.id)
        out.append(
            ProductPriceListRef(
                price_list_id=pl.id,
                price_list_name=pl.name,
                unit_price=Decimal(item.unit_price or 0),
                is_active=bool(pl.is_active),
                valid_from=pl.valid_from,
                valid_to=pl.valid_to,
            )
        )
    return out

