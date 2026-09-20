"""Fiyat listeleri CRUD + yazdır / CSV / PDF."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from decimal import Decimal

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session, joinedload

from app.core.deps import get_db, require_roles
from app.models.price_list import PriceList, PriceListItem
from app.models.user import User
from app.schemas.price_list import (
    PriceListCreate,
    PriceListListItem,
    PriceListOut,
    PriceListUpdate,
)
from app.services.audit import write_audit

router = APIRouter(prefix="/price-lists", tags=["price-lists"])


def _d(v) -> Decimal:
    return Decimal("0") if v is None else Decimal(str(v))


def _to_list(pl: PriceList) -> PriceListListItem:
    return PriceListListItem(
        id=pl.id,
        name=pl.name,
        description=pl.description,
        currency=pl.currency or "TRY",
        is_active=bool(pl.is_active),
        valid_from=pl.valid_from,
        valid_to=pl.valid_to,
        item_count=len(pl.items or []),
        created_at=pl.created_at,
        updated_at=pl.updated_at,
    )


def _to_out(pl: PriceList) -> PriceListOut:
    return PriceListOut(**_to_list(pl).model_dump(), items=pl.items or [])


def _load(db: Session, list_id: int) -> PriceList:
    pl = (
        db.query(PriceList)
        .options(joinedload(PriceList.items))
        .filter(PriceList.id == list_id)
        .first()
    )
    if not pl:
        raise HTTPException(status_code=404, detail="Fiyat listesi bulunamadı")
    return pl


def _replace_items(pl: PriceList, items: list) -> None:
    pl.items.clear()
    for item in items:
        data = item.model_dump() if hasattr(item, "model_dump") else dict(item)
        blank = data.get("blank_price")
        unit = data.get("unit_price")
        # Baskısız yoksa unit_price kullan; unit_price yoksa baskısız
        if blank is None and unit is not None:
            blank = unit
        if unit is None or (blank is not None and _d(unit) == 0 and _d(blank) > 0):
            unit = blank if blank is not None else 0
        pl.items.append(
            PriceListItem(
                product_id=data.get("product_id"),
                variant_id=data.get("variant_id"),
                description=data["description"],
                unit_price=_d(unit),
                supplier_name=(data.get("supplier_name") or None),
                purchase_price=_d(data["purchase_price"]) if data.get("purchase_price") is not None else None,
                blank_price=_d(blank) if blank is not None else None,
                printed_price=_d(data["printed_price"]) if data.get("printed_price") is not None else None,
                embroidered_price=_d(data["embroidered_price"]) if data.get("embroidered_price") is not None else None,
                valid_from=data.get("valid_from"),
                valid_to=data.get("valid_to"),
                notes=data.get("notes"),
            )
        )


@router.get("", response_model=list[PriceListListItem])
def list_price_lists(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "satış", "muhasebe")),
    q: str | None = Query(default=None),
    active_only: bool = False,
) -> list[PriceListListItem]:
    query = db.query(PriceList).options(joinedload(PriceList.items))
    if active_only:
        query = query.filter(PriceList.is_active.is_(True))
    if q:
        like = f"%{q}%"
        query = query.filter((PriceList.name.ilike(like)) | (PriceList.description.ilike(like)))
    rows = query.order_by(PriceList.id.desc()).all()
    return [_to_list(r) for r in rows]


@router.post("", response_model=PriceListOut, status_code=status.HTTP_201_CREATED)
def create_price_list(
    payload: PriceListCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles("admin", "satış")),
) -> PriceListOut:
    pl = PriceList(
        name=payload.name.strip(),
        description=payload.description,
        currency=payload.currency or "TRY",
        is_active=payload.is_active,
        valid_from=payload.valid_from,
        valid_to=payload.valid_to,
    )
    db.add(pl)
    db.flush()
    _replace_items(pl, payload.items)
    db.commit()
    out = _to_out(_load(db, pl.id))
    write_audit(
        user_id=user.id,
        action="create",
        entity_type="price_list",
        entity_id=pl.id,
        detail={"name": pl.name},
    )
    return out



PRICE_LIST_TEMPLATE_HEADERS = [
    "product_id",
    "SKU",
    "Barkod",
    "Ürün",
    "Tedarikçi",
    "Alış Fiyatı",
    "Baskısız Fiyatı",
    "Baskılı Fiyatı",
    "Nakışlı Fiyatı",
    "Not",
]

PRICE_FIELDS = ("purchase_price", "blank_price", "printed_price", "embroidered_price")


class BulkAdjustRequest(BaseModel):
    """Seçili kalemlerde % veya mutlak fiyat güncelleme."""

    item_ids: list[int] = Field(default_factory=list)
    mode: str = Field(default="percent", pattern="^(percent|absolute)$")
    value: float = 0
    fields: list[str] = Field(default_factory=lambda: ["blank_price", "printed_price", "embroidered_price", "purchase_price"])
    apply_all: bool = False


class BulkAdjustResult(BaseModel):
    updated: int
    skipped: int


def _norm_header(h: str) -> str:
    return str(h or "").strip().lower().replace("ı", "i").replace("İ", "i")


def _col_map(headers: list[str]) -> dict[str, str]:
    clean = {_norm_header(h): h for h in headers}
    adaylar = {
        "product_id": ["product_id", "urun_id", "ürün_id", "urun id", "ürün id", "id"],
        "sku": ["sku", "stok kodu", "stok_kodu", "kod"],
        "barcode": ["barkod", "barcode", "ean", "barcode_no"],
        "description": ["urun", "ürün", "urun adi", "ürün adı", "product", "product_name", "name"],
        "supplier_name": ["tedarikci", "tedarikçi", "supplier"],
        "purchase_price": ["alis fiyati", "alış fiyatı", "alis", "alış", "maliyet", "purchase"],
        "blank_price": ["baskisiz fiyati", "baskısız fiyatı", "baskisiz", "baskısız", "blank", "unit_price", "satis", "satış"],
        "printed_price": ["baskili fiyati", "baskılı fiyatı", "baskili", "baskılı", "printed"],
        "embroidered_price": ["nakisli fiyati", "nakışlı fiyatı", "nakisli", "nakışlı", "embroidered"],
        "notes": ["not", "notes"],
    }
    out: dict[str, str] = {}
    for hedef, isimler in adaylar.items():
        for ad in isimler:
            if ad in clean:
                out[hedef] = clean[ad]
                break
    return out


@router.get("/import-template")
def price_list_import_template(
    _: User = Depends(require_roles("admin", "satış", "muhasebe")),
    fmt: str = Query(default="xlsx", pattern="^(csv|xlsx)$"),
):
    """Masaüstü FIYAT_LISTESI_KOLONLARI — boş fiyat listesi şablonu."""
    headers = PRICE_LIST_TEMPLATE_HEADERS
    if fmt == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="openpyxl yüklü değil") from exc
        wb = Workbook()
        ws = wb.active
        ws.title = "FiyatListesi"
        ws.append(headers)
        buf = io.BytesIO()
        wb.save(buf)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="fiyat-listesi-sablon.xlsx"'},
        )
    buf = io.StringIO()
    csv.writer(buf).writerow(headers)
    return Response(
        content=buf.getvalue().encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="fiyat-listesi-sablon.csv"'},
    )


@router.post("/{list_id}/bulk-adjust", response_model=BulkAdjustResult)
def bulk_adjust_prices(
    list_id: int,
    payload: BulkAdjustRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles("admin", "satış")),
) -> BulkAdjustResult:
    pl = _load(db, list_id)
    fields = [f for f in payload.fields if f in PRICE_FIELDS]
    if not fields:
        raise HTTPException(status_code=400, detail="En az bir fiyat alanı seçin")
    id_set = set(payload.item_ids or [])
    updated = 0
    skipped = 0
    for it in pl.items or []:
        if not payload.apply_all and it.id not in id_set:
            continue
        touched = False
        for f in fields:
            cur = getattr(it, f, None)
            if cur is None and f == "blank_price":
                cur = it.unit_price
            if cur is None:
                cur = Decimal("0")
            cur_d = _d(cur)
            if payload.mode == "percent":
                new_v = cur_d * (Decimal("1") + Decimal(str(payload.value)) / Decimal("100"))
            else:
                new_v = cur_d + Decimal(str(payload.value))
            if new_v < 0:
                new_v = Decimal("0")
            setattr(it, f, new_v)
            if f == "blank_price":
                it.unit_price = new_v
            touched = True
        if touched:
            updated += 1
        else:
            skipped += 1
    pl.updated_at = datetime.utcnow()
    db.commit()
    write_audit(
        user_id=user.id,
        action="bulk_adjust",
        entity_type="price_list",
        entity_id=pl.id,
        detail={"mode": payload.mode, "value": payload.value, "updated": updated, "fields": fields},
    )
    return BulkAdjustResult(updated=updated, skipped=skipped)


@router.post("/{list_id}/import")
async def import_price_list_rows(
    list_id: int,
    file: UploadFile = File(...),
    mode: str = Query(default="merge", pattern="^(merge|replace)$"),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles("admin", "satış")),
) -> dict:
    """Excel/CSV satırlarını fiyat listesine aktar (merge: eşleşen ürün güncelle / yeni ekle)."""
    pl = _load(db, list_id)
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="Boş dosya")
    name = (file.filename or "").lower()
    rows: list[dict] = []
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="openpyxl yüklü değil") from exc
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        if not data:
            raise HTTPException(status_code=400, detail="Boş dosya")
        headers = [str(h or "").strip() for h in data[0]]
        for line in data[1:]:
            if not line or all(v is None or str(v).strip() == "" for v in line):
                continue
            rows.append({headers[i]: line[i] if i < len(line) else None for i in range(len(headers))})
    else:
        text_data = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text_data))
        rows = [r for r in reader if any(str(v or "").strip() for v in r.values())]

    if not rows:
        raise HTTPException(status_code=400, detail="Excel dosyasında okunacak satır bulunamadı")

    cmap = _col_map(list(rows[0].keys()))
    has_id_keys = any(k in cmap for k in ("product_id", "sku", "barcode"))
    if "description" not in cmap and not has_id_keys:
        raise HTTPException(
            status_code=400,
            detail="Excel içinde product_id / SKU / Barkod veya Ürün sütunu olmalı.",
        )

    def _num(v) -> Decimal | None:
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "nan":
            return None
        try:
            s = str(v).strip().replace(",", ".")
            return Decimal(s)
        except Exception:
            return None

    from app.models.product import Product, ProductVariant

    products = db.query(Product).all()
    by_id: dict[int, Product] = {p.id: p for p in products}
    by_sku: dict[str, Product] = {str(p.sku or "").strip().lower(): p for p in products if p.sku}
    by_name: dict[str, Product] = {str(p.name or "").strip().lower(): p for p in products if p.name}
    by_barcode: dict[str, Product] = {}
    for v in db.query(ProductVariant).all():
        bc = str(v.barcode or "").strip()
        if bc and v.product_id in by_id:
            by_barcode[bc.lower()] = by_id[v.product_id]
        sku = str(v.sku or "").strip().lower()
        if sku and v.product_id in by_id and sku not in by_sku:
            by_sku[sku] = by_id[v.product_id]

    def _resolve_product(row: dict) -> tuple[Product | None, str]:
        """Match product_id → barcode → SKU → name. Returns (product, match_how)."""
        if "product_id" in cmap:
            raw = row.get(cmap["product_id"])
            if raw is not None and str(raw).strip() and str(raw).strip().lower() != "nan":
                try:
                    pid = int(float(str(raw).strip()))
                    if pid in by_id:
                        return by_id[pid], "product_id"
                    return None, f"product_id={pid} bulunamadı"
                except (TypeError, ValueError):
                    return None, f"product_id geçersiz ({raw})"
        if "barcode" in cmap:
            bc = str(row.get(cmap["barcode"]) or "").strip()
            if bc and bc.lower() != "nan":
                p = by_barcode.get(bc.lower())
                if p:
                    return p, "barcode"
                return None, f"barkod={bc} bulunamadı"
        if "sku" in cmap:
            sku = str(row.get(cmap["sku"]) or "").strip()
            if sku and sku.lower() != "nan":
                p = by_sku.get(sku.lower())
                if p:
                    return p, "sku"
                return None, f"sku={sku} bulunamadı"
        if "description" in cmap:
            name = str(row.get(cmap["description"]) or "").strip()
            if name and name.lower() != "nan":
                p = by_name.get(name.lower())
                if p:
                    return p, "name"
                return None, f"isim={name} eşleşmedi"
        return None, "eşleştirme anahtarı yok"

    created = 0
    updated = 0
    skipped = 0
    unmatched = 0
    errors: list[str] = []
    unmatched_rows: list[dict] = []

    if mode == "replace":
        pl.items.clear()
        db.flush()

    existing_by_pid: dict[int, PriceListItem] = {}
    existing_by_desc: dict[str, PriceListItem] = {}
    for it in pl.items or []:
        if it.product_id:
            existing_by_pid[int(it.product_id)] = it
        existing_by_desc[str(it.description or "").strip().lower()] = it

    for idx, row in enumerate(rows, start=2):
        prod, how = _resolve_product(row)
        desc_from_row = ""
        if "description" in cmap:
            desc_from_row = str(row.get(cmap["description"]) or "").strip()
            if desc_from_row.lower() == "nan":
                desc_from_row = ""

        if prod is None:
            # Free-text rows: need a description; report catalog miss clearly
            if not desc_from_row:
                skipped += 1
                unmatched += 1
                errors.append(f"Satır {idx}: eşleşmedi ({how})")
                unmatched_rows.append({"row": idx, "reason": how})
                continue
            desc = desc_from_row
            product_id = None
            unmatched += 1
            errors.append(f"Satır {idx}: ürün kataloğunda yok — isimle eklendi/güncellendi ({how})")
            unmatched_rows.append({"row": idx, "reason": how, "description": desc})
        else:
            desc = desc_from_row or prod.name
            product_id = prod.id

        supplier = str(row.get(cmap.get("supplier_name", ""), "") or "").strip() or None
        purchase = _num(row.get(cmap["purchase_price"])) if "purchase_price" in cmap else None
        blank = _num(row.get(cmap["blank_price"])) if "blank_price" in cmap else None
        printed = _num(row.get(cmap["printed_price"])) if "printed_price" in cmap else None
        emb = _num(row.get(cmap["embroidered_price"])) if "embroidered_price" in cmap else None
        notes = str(row.get(cmap.get("notes", ""), "") or "").strip() or None
        if purchase is None and blank is None and printed is None and emb is None:
            skipped += 1
            errors.append(f"Satır {idx}: fiyat yok — atlandı ({desc})")
            continue

        it = None
        if product_id is not None:
            it = existing_by_pid.get(product_id)
        if it is None:
            it = existing_by_desc.get(desc.lower())

        if it is None:
            unit = blank if blank is not None else Decimal("0")
            it = PriceListItem(
                product_id=product_id,
                description=desc,
                supplier_name=supplier,
                purchase_price=purchase,
                blank_price=blank,
                printed_price=printed,
                embroidered_price=emb,
                unit_price=unit,
                notes=notes,
            )
            pl.items.append(it)
            if product_id is not None:
                existing_by_pid[product_id] = it
            existing_by_desc[desc.lower()] = it
            created += 1
        else:
            if product_id is not None:
                it.product_id = product_id
            it.description = desc
            if supplier is not None:
                it.supplier_name = supplier
            if purchase is not None:
                it.purchase_price = purchase
            if blank is not None:
                it.blank_price = blank
                it.unit_price = blank
            if printed is not None:
                it.printed_price = printed
            if emb is not None:
                it.embroidered_price = emb
            if notes is not None:
                it.notes = notes
            updated += 1

    pl.updated_at = datetime.utcnow()
    db.commit()
    write_audit(
        user_id=user.id,
        action="import",
        entity_type="price_list",
        entity_id=pl.id,
        detail={
            "created": created,
            "updated": updated,
            "skipped": skipped,
            "unmatched": unmatched,
            "mode": mode,
        },
    )
    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "unmatched": unmatched,
        "unmatched_rows": unmatched_rows[:50],
        "errors": errors[:40],
        "message": (
            f"İçe aktarma: {created} yeni, {updated} güncellendi, {skipped} atlandı"
            + (f", {unmatched} eşleşmeyen" if unmatched else "")
        ),
    }


@router.get("/{list_id}", response_model=PriceListOut)
def get_price_list(
    list_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "satış", "muhasebe")),
) -> PriceListOut:
    return _to_out(_load(db, list_id))


@router.put("/{list_id}", response_model=PriceListOut)
def update_price_list(
    list_id: int,
    payload: PriceListUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles("admin", "satış")),
) -> PriceListOut:
    pl = _load(db, list_id)
    data = payload.model_dump(exclude_unset=True)
    items = data.pop("items", None)
    for k, v in data.items():
        setattr(pl, k, v)
    if items is not None:
        class Wrap:
            def __init__(self, d: dict):
                self._d = d

            def model_dump(self):
                return self._d

        _replace_items(pl, [Wrap(x) for x in items])
    pl.updated_at = datetime.utcnow()
    db.commit()
    write_audit(
        user_id=user.id,
        action="update",
        entity_type="price_list",
        entity_id=pl.id,
        detail={"name": pl.name},
    )
    return _to_out(_load(db, pl.id))


@router.delete("/{list_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_price_list(
    list_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles("admin")),
) -> None:
    pl = _load(db, list_id)
    name = pl.name
    db.delete(pl)
    db.commit()
    write_audit(
        user_id=user.id,
        action="delete",
        entity_type="price_list",
        entity_id=list_id,
        detail={"name": name},
    )


def _item_rows(pl: PriceList) -> list[dict]:
    rows = []
    for it in pl.items or []:
        blank = it.blank_price if it.blank_price is not None else it.unit_price
        rows.append(
            {
                "description": it.description,
                "supplier_name": it.supplier_name or "",
                "purchase_price": float(it.purchase_price or 0),
                "blank_price": float(blank or 0),
                "printed_price": float(it.printed_price or 0),
                "embroidered_price": float(it.embroidered_price or 0),
                "unit_price": float(it.unit_price or 0),
                "notes": it.notes or "",
            }
        )
    return rows


@router.get("/{list_id}/export")
def export_price_list(
    list_id: int,
    fmt: str = Query(default="csv", pattern="^(csv|html|pdf)$"),
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "satış", "muhasebe")),
):
    """CSV / yazdırılabilir HTML / basit PDF."""
    pl = _load(db, list_id)
    rows = _item_rows(pl)
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(
            ["Ürün", "Tedarikçi", "Alış", "Baskısız", "Baskılı", "Nakışlı", "Not"]
        )
        for r in rows:
            w.writerow(
                [
                    r["description"],
                    r["supplier_name"],
                    r["purchase_price"],
                    r["blank_price"],
                    r["printed_price"],
                    r["embroidered_price"],
                    r["notes"],
                ]
            )
        data = "\ufeff" + buf.getvalue()
        return StreamingResponse(
            iter([data]),
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f'attachment; filename="fiyat_listesi_{list_id}.csv"'
            },
        )
    if fmt == "html":
        trs = "".join(
            f"<tr><td>{r['description']}</td><td>{r['supplier_name']}</td>"
            f"<td style='text-align:right'>{r['purchase_price']:.2f}</td>"
            f"<td style='text-align:right'>{r['blank_price']:.2f}</td>"
            f"<td style='text-align:right'>{r['printed_price']:.2f}</td>"
            f"<td style='text-align:right'>{r['embroidered_price']:.2f}</td>"
            f"<td>{r['notes']}</td></tr>"
            for r in rows
        )
        html = f"""<!DOCTYPE html><html><head><meta charset="utf-8">
<title>{pl.name}</title>
<style>
body{{font-family:Segoe UI,Arial,sans-serif;padding:24px;color:#0f172a}}
h1{{font-size:20px;margin:0 0 4px}} .muted{{color:#64748b;font-size:12px;margin-bottom:16px}}
table{{border-collapse:collapse;width:100%;font-size:12px}}
th,td{{border:1px solid #cbd5e1;padding:6px 8px}} th{{background:#0f766e;color:#fff;text-align:left}}
@media print{{button{{display:none}}}}
</style></head><body>
<button onclick="window.print()">Yazdır</button>
<h1>{pl.name}</h1>
<div class="muted">Fiyat / Maliyet › Fiyat Listesi · {len(rows)} kalem</div>
<table><thead><tr>
<th>Ürün</th><th>Tedarikçi</th><th>Alış</th><th>Baskısız</th><th>Baskılı</th><th>Nakışlı</th><th>Not</th>
</tr></thead><tbody>{trs or '<tr><td colspan=7>Kalem yok</td></tr>'}</tbody></table>
</body></html>"""
        return HTMLResponse(html)
    # pdf
    from app.models.settings_model import AppSetting
    from app.services.pdf import build_price_list_pdf

    settings_map = {s.key: (s.value or "") for s in db.query(AppSetting).all()}
    pdf = build_price_list_pdf(pl.name, rows, settings_map)
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="fiyat_listesi_{list_id}.pdf"'},
    )
