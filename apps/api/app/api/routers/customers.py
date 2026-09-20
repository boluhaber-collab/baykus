"""Customers CRUD + Cari (current account) ledger."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload

from app.core.deps import get_db, require_roles
from app.services.audit import write_audit
from app.models.customer import CARI_MOVEMENT_TYPES, CariMovement, Customer
from app.models.order import CLOSED_STATUSES, Order, Payment
from app.models.quote import Quote
from app.models.user import User
from app.schemas.customer import (
    CariMovementCreate,
    CariMovementOut,
    CustomerCreate,
    CustomerDetailOut,
    CustomerDevirIn,
    CustomerOrderBrief,
    CustomerQuoteBrief,
    CustomerOut,
    CustomerTahsilatIn,
    CustomerTahsilatOut,
    CustomerUpdate,
    ReceivableItem,
    StatementOut,
)

router = APIRouter(prefix="/customers", tags=["customers"])

READ_ROLES = ("admin", "satış", "muhasebe")
WRITE_ROLES = ("admin", "satış")
CARI_WRITE_ROLES = ("admin", "satış", "muhasebe")


def _balance_for(db: Session, customer: Customer) -> Decimal:
    row = (
        db.query(
            func.coalesce(func.sum(CariMovement.debit), 0),
            func.coalesce(func.sum(CariMovement.credit), 0),
        )
        .filter(CariMovement.customer_id == customer.id)
        .one()
    )
    debit_sum = Decimal(str(row[0]))
    credit_sum = Decimal(str(row[1]))
    opening = customer.opening_balance or Decimal("0")
    return opening + debit_sum - credit_sum


def _balances_map(db: Session, customer_ids: list[int]) -> dict[int, tuple[Decimal, Decimal]]:
    """Return {customer_id: (debit_sum, credit_sum)}."""
    if not customer_ids:
        return {}
    rows = (
        db.query(
            CariMovement.customer_id,
            func.coalesce(func.sum(CariMovement.debit), 0),
            func.coalesce(func.sum(CariMovement.credit), 0),
        )
        .filter(CariMovement.customer_id.in_(customer_ids))
        .group_by(CariMovement.customer_id)
        .all()
    )
    return {r[0]: (Decimal(str(r[1])), Decimal(str(r[2]))) for r in rows}


def _customer_out(customer: Customer, balance: Decimal) -> CustomerOut:
    return CustomerOut(
        id=customer.id,
        code=customer.code,
        name=customer.name,
        company=customer.company,
        email=customer.email,
        phone=customer.phone,
        city=customer.city,
        address=customer.address,
        tax_number=customer.tax_number,
        tax_office=customer.tax_office,
        notes=customer.notes,
        is_active=bool(customer.is_active),
        special_day_note=customer.special_day_note,
        special_day_date=customer.special_day_date,
        opening_balance=customer.opening_balance or Decimal("0"),
        balance=balance,
        created_at=customer.created_at,
        updated_at=customer.updated_at,
    )


def _movement_out(m: CariMovement, running: Decimal | None = None) -> CariMovementOut:
    order_number = None
    if m.order is not None:
        order_number = m.order.order_number
    return CariMovementOut(
        id=m.id,
        customer_id=m.customer_id,
        movement_type=m.movement_type,
        debit=m.debit or Decimal("0"),
        credit=m.credit or Decimal("0"),
        movement_date=m.movement_date,
        order_id=m.order_id,
        order_number=order_number,
        note=m.note,
        created_at=m.created_at,
        running_balance=running,
    )


def _split_amount(movement_type: str, amount: Decimal, side: str | None) -> tuple[Decimal, Decimal]:
    if movement_type == "sale":
        return amount, Decimal("0")
    if movement_type in ("payment", "deposit"):
        return Decimal("0"), amount
    # adjustment
    if side == "credit":
        return Decimal("0"), amount
    return amount, Decimal("0")


def _paid_for_order(order: Order) -> Decimal:
    return sum((p.amount for p in order.payments), Decimal("0"))


@router.get("", response_model=list[CustomerOut])
def list_customers(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
    q: str | None = Query(default=None),
    active: bool | None = Query(default=None),
    has_balance: bool | None = Query(default=None),
    skip: int = 0,
    limit: int = 200,
) -> list[CustomerOut]:
    query = db.query(Customer)
    if q:
        like = f"%{q}%"
        query = query.filter(
            or_(
                Customer.name.ilike(like),
                Customer.company.ilike(like),
                Customer.email.ilike(like),
                Customer.phone.ilike(like),
                Customer.code.ilike(like),
                Customer.tax_number.ilike(like),
            )
        )
    if active is not None:
        query = query.filter(Customer.is_active.is_(active))
    customers = query.order_by(Customer.name.asc()).offset(skip).limit(limit).all()
    bal_map = _balances_map(db, [c.id for c in customers])
    result: list[CustomerOut] = []
    for c in customers:
        debit_s, credit_s = bal_map.get(c.id, (Decimal("0"), Decimal("0")))
        balance = (c.opening_balance or Decimal("0")) + debit_s - credit_s
        if has_balance is True and balance <= 0:
            continue
        if has_balance is False and balance > 0:
            continue
        result.append(_customer_out(c, balance))
    return result




@router.get("/track")
def customer_track(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
    limit: int = Query(default=50, ge=1, le=200),
    customer_id: int | None = Query(default=None),
) -> dict:
    """Müşteri takip — timeline + açık alacak özeti + notlar (masaüstü)."""
    events: list[dict] = []
    oq = db.query(Order).options(joinedload(Order.customer)).order_by(Order.created_at.desc())
    if customer_id:
        oq = oq.filter(Order.customer_id == customer_id)
    orders = oq.limit(limit).all()
    for o in orders:
        events.append(
            {
                "type": "order",
                "at": o.created_at.isoformat() if o.created_at else None,
                "customer_id": o.customer_id,
                "customer_name": o.customer.name if o.customer else None,
                "label": f"Sipariş {o.order_number}",
                "detail": o.status,
                "amount": float(o.total_amount or 0),
                "href": f"/orders/{o.id}",
            }
        )
    mq = (
        db.query(CariMovement)
        .options(joinedload(CariMovement.customer))
        .order_by(CariMovement.movement_date.desc(), CariMovement.id.desc())
    )
    if customer_id:
        mq = mq.filter(CariMovement.customer_id == customer_id)
    else:
        mq = mq.filter(CariMovement.movement_type == "payment")
    payments = mq.limit(limit).all()
    for m in payments:
        events.append(
            {
                "type": m.movement_type or "payment",
                "at": m.movement_date.isoformat() if m.movement_date else None,
                "customer_id": m.customer_id,
                "customer_name": m.customer.name if m.customer else None,
                "label": "Cari " + (m.movement_type or "hareket"),
                "detail": m.note,
                "amount": float(m.credit or m.debit or 0),
                "href": f"/customers/{m.customer_id}",
            }
        )
    events.sort(key=lambda e: e.get("at") or "", reverse=True)

    # Receivables snapshot
    receivables = []
    customers_q = db.query(Customer).filter(Customer.is_active.is_(True))
    if customer_id:
        customers_q = customers_q.filter(Customer.id == customer_id)
    customers = customers_q.order_by(Customer.name).all()
    bal_map = _balances_map(db, [c.id for c in customers])
    for c in customers:
        debit_s, credit_s = bal_map.get(c.id, (Decimal("0"), Decimal("0")))
        balance = (c.opening_balance or Decimal("0")) + debit_s - credit_s
        if balance <= 0 and not customer_id:
            continue
        receivables.append(
            {
                "customer_id": c.id,
                "name": c.name,
                "phone": c.phone,
                "balance": float(balance),
                "notes": c.notes,
                "special_day_note": c.special_day_note,
                "special_day_date": c.special_day_date.isoformat() if c.special_day_date else None,
            }
        )
    receivables.sort(key=lambda x: x["balance"], reverse=True)

    focus = None
    if customer_id:
        c = db.get(Customer, customer_id)
        if c:
            debit_s, credit_s = bal_map.get(c.id, (Decimal("0"), Decimal("0")))
            balance = (c.opening_balance or Decimal("0")) + debit_s - credit_s
            focus = {
                "id": c.id,
                "name": c.name,
                "phone": c.phone,
                "company": c.company,
                "notes": c.notes,
                "special_day_note": c.special_day_note,
                "special_day_date": c.special_day_date.isoformat() if c.special_day_date else None,
                "balance": float(balance),
            }

    return {
        "events": events[:limit],
        "receivables": receivables[:40],
        "focus": focus,
        "summary": {
            "event_count": len(events[:limit]),
            "receivable_count": len(receivables),
            "receivable_total": sum(r["balance"] for r in receivables if r["balance"] > 0),
        },
    }

@router.get("/receivables", response_model=list[ReceivableItem])
def open_receivables(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
) -> list[ReceivableItem]:
    customers = db.query(Customer).filter(Customer.is_active.is_(True)).order_by(Customer.name).all()
    bal_map = _balances_map(db, [c.id for c in customers])
    last_dates = dict(
        db.query(CariMovement.customer_id, func.max(CariMovement.movement_date))
        .group_by(CariMovement.customer_id)
        .all()
    )
    items: list[ReceivableItem] = []
    for c in customers:
        debit_s, credit_s = bal_map.get(c.id, (Decimal("0"), Decimal("0")))
        balance = (c.opening_balance or Decimal("0")) + debit_s - credit_s
        if balance <= 0:
            continue
        items.append(
            ReceivableItem(
                customer_id=c.id,
                code=c.code,
                name=c.name,
                company=c.company,
                phone=c.phone,
                city=c.city,
                balance=balance,
                last_movement_date=last_dates.get(c.id),
            )
        )
    items.sort(key=lambda x: x.balance, reverse=True)
    return items


@router.post("", response_model=CustomerOut, status_code=status.HTTP_201_CREATED)
def create_customer(
    payload: CustomerCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*WRITE_ROLES)),
) -> CustomerOut:
    data = payload.model_dump()
    if data.get("code"):
        existing = db.query(Customer).filter(Customer.code == data["code"]).first()
        if existing:
            raise HTTPException(status_code=400, detail="Müşteri kodu zaten kullanılıyor")
    customer = Customer(**data)
    db.add(customer)
    db.commit()
    db.refresh(customer)
    write_audit(
        user_id=user.id,
        action="create",
        entity_type="customer",
        entity_id=customer.id,
        detail={"name": customer.name},
    )
    return _customer_out(customer, customer.opening_balance or Decimal("0"))


@router.get("/import-template")
def customer_import_template(
    _: User = Depends(require_roles(*WRITE_ROLES)),
    fmt: str = Query(default="csv", pattern="^(csv|xlsx)$"),
):
    """Excelden Müşteri — boş şablon (adı zorunlu)."""
    import csv
    import io

    headers = [
        "Ad",
        "Firma",
        "Telefon",
        "E-posta",
        "Şehir",
        "Adres",
        "Vergi No",
        "Vergi Dairesi",
        "Kod",
        "Not",
    ]
    if fmt == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="openpyxl yüklü değil") from exc
        wb = Workbook()
        ws = wb.active
        ws.title = "Musteriler"
        ws.append(headers)
        buf = io.BytesIO()
        wb.save(buf)
        return Response(
            buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="musteri-sablon.xlsx"'},
        )
    buf = io.StringIO()
    csv.writer(buf).writerow(headers)
    return Response(
        buf.getvalue().encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="musteri-sablon.csv"'},
    )


@router.post("/import")
async def import_customers(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*WRITE_ROLES)),
) -> dict:
    """Excel/CSV müşteri içe aktarma — şifre/sır yok."""
    import csv
    import io

    raw = await file.read()
    name = (file.filename or "").lower()
    rows: list[dict] = []
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="openpyxl yüklü değil") from exc
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        if not data:
            raise HTTPException(status_code=400, detail="Boş dosya")
        headers = [str(h or "").strip() for h in data[0]]
        for line in data[1:]:
            rows.append({headers[i]: line[i] if i < len(line) else None for i in range(len(headers))})
    else:
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)

    def cell(row: dict, *keys: str) -> str:
        lower = {str(k).strip().casefold(): v for k, v in row.items() if k is not None}
        aliases = {
            "ad": ("ad", "adı", "adi", "name", "müşteri", "musteri", "ünvan", "unvan"),
            "firma": ("firma", "şirket", "sirket", "company"),
            "telefon": ("telefon", "tel", "phone", "gsm", "cep"),
            "email": ("e-posta", "eposta", "email", "mail"),
            "şehir": ("şehir", "sehir", "city"),
            "adres": ("adres", "address"),
            "vergi_no": ("vergi no", "vergi_no", "vkn", "tax_number", "tc"),
            "vergi_dairesi": ("vergi dairesi", "vergi_dairesi", "tax_office"),
            "kod": ("kod", "code", "cari kod", "müşteri kodu"),
            "not": ("not", "notes", "açıklama", "aciklama"),
        }
        for canon in keys:
            for alias in aliases.get(canon, (canon,)):
                if alias.casefold() in lower:
                    v = lower[alias.casefold()]
                    if v is None:
                        return ""
                    return str(v).strip()
        return ""

    created = updated = skipped = 0
    errors: list[str] = []
    for idx, row in enumerate(rows, start=2):
        ad = cell(row, "ad")
        if not ad or ad.casefold() in ("nan", "none"):
            skipped += 1
            continue
        telefon = cell(row, "telefon")
        email = cell(row, "email") or None
        if email and "@" not in email:
            email = None
        firma = cell(row, "firma") or None
        sehir = cell(row, "şehir") or None
        adres = cell(row, "adres") or None
        vergi_no = cell(row, "vergi_no") or None
        vergi_d = cell(row, "vergi_dairesi") or None
        kod = cell(row, "kod") or None
        notu = cell(row, "not") or None

        existing = None
        if telefon:
            existing = (
                db.query(Customer)
                .filter(Customer.phone == telefon)
                .first()
            )
        if not existing and kod:
            existing = db.query(Customer).filter(Customer.code == kod).first()
        if not existing:
            # name+phone soft match
            q = db.query(Customer).filter(Customer.name == ad)
            if telefon:
                q = q.filter(Customer.phone == telefon)
            existing = q.first()

        if existing:
            existing.name = ad
            if firma:
                existing.company = firma
            if telefon:
                existing.phone = telefon
            if email:
                existing.email = email
            if sehir:
                existing.city = sehir
            if adres:
                existing.address = adres
            if vergi_no:
                existing.tax_number = vergi_no
            if vergi_d:
                existing.tax_office = vergi_d
            if kod and not existing.code:
                existing.code = kod
            if notu:
                existing.notes = ((existing.notes or "") + " | " + notu).strip(" |")
            updated += 1
        else:
            db.add(
                Customer(
                    name=ad,
                    company=firma,
                    phone=telefon or None,
                    email=email,
                    city=sehir,
                    address=adres,
                    tax_number=vergi_no,
                    tax_office=vergi_d,
                    code=kod,
                    notes=notu,
                    is_active=True,
                )
            )
            created += 1

    db.commit()
    write_audit(
        user_id=user.id if user else None,
        action="customer_import",
        entity_type="customers",
        detail=f"created={created} updated={updated} skipped={skipped}",
    )
    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:20],
        "message": f"Eklenen: {created} · Güncellenen: {updated} · Atlanan: {skipped}",
    }


@router.get("/{customer_id}", response_model=CustomerDetailOut)
def get_customer_detail(
    customer_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
) -> CustomerDetailOut:
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Müşteri bulunamadı")
    balance = _balance_for(db, customer)
    base = _customer_out(customer, balance)

    orders = (
        db.query(Order)
        .filter(Order.customer_id == customer_id)
        .order_by(Order.created_at.desc())
        .limit(10)
        .all()
    )
    recent_orders: list[CustomerOrderBrief] = []
    for o in orders:
        paid = _paid_for_order(o)
        remaining = (o.total_amount or Decimal("0")) - paid
        if remaining < 0:
            remaining = Decimal("0")
        recent_orders.append(
            CustomerOrderBrief(
                id=o.id,
                order_number=o.order_number,
                status=o.status,
                total_amount=o.total_amount or Decimal("0"),
                remaining_amount=remaining,
                due_date=o.due_date,
                created_at=o.created_at,
            )
        )


    quotes = (
        db.query(Quote)
        .filter(Quote.customer_id == customer_id)
        .order_by(Quote.created_at.desc())
        .limit(10)
        .all()
    )
    recent_quotes = [
        CustomerQuoteBrief(
            id=q.id,
            quote_number=q.quote_number,
            status=q.status,
            total_amount=q.total_amount or Decimal("0"),
            valid_until=q.valid_until,
            created_at=q.created_at,
        )
        for q in quotes
    ]
    movements = (
        db.query(CariMovement)
        .filter(CariMovement.customer_id == customer_id)
        .order_by(CariMovement.movement_date.desc(), CariMovement.id.desc())
        .limit(15)
        .all()
    )
    recent_movements = [_movement_out(m) for m in movements]

    timeline: list[dict] = []
    for o in recent_orders:
        timeline.append(
            {
                "kind": "order",
                "date": o.created_at.date().isoformat() if o.created_at else None,
                "label": f"Sipariş {o.order_number}",
                "status": o.status,
                "amount": float(o.total_amount),
                "ref_id": o.id,
            }
        )
    for q in recent_quotes:
        timeline.append(
            {
                "kind": "quote",
                "date": q.created_at.date().isoformat() if q.created_at else None,
                "label": f"Teklif {q.quote_number}",
                "status": q.status,
                "amount": float(q.total_amount),
                "ref_id": q.id,
            }
        )
    for m in recent_movements:
        timeline.append(
            {
                "kind": "movement",
                "date": m.movement_date.isoformat(),
                "label": m.movement_type,
                "note": m.note,
                "debit": float(m.debit),
                "credit": float(m.credit),
                "ref_id": m.id,
            }
        )
    if customer.notes:
        timeline.append(
            {
                "kind": "note",
                "date": customer.updated_at.date().isoformat() if customer.updated_at else None,
                "label": "Müşteri notu",
                "note": customer.notes,
            }
        )
    timeline.sort(key=lambda x: x.get("date") or "", reverse=True)

    return CustomerDetailOut(
        **base.model_dump(),
        recent_orders=recent_orders,
        recent_quotes=recent_quotes,
        recent_movements=recent_movements,
        timeline=timeline[:20],
    )


@router.put("/{customer_id}", response_model=CustomerOut)
def update_customer(
    customer_id: int,
    payload: CustomerUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*WRITE_ROLES)),
) -> CustomerOut:
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Müşteri bulunamadı")
    data = payload.model_dump(exclude_unset=True)
    if "code" in data and data["code"]:
        clash = (
            db.query(Customer)
            .filter(Customer.code == data["code"], Customer.id != customer_id)
            .first()
        )
        if clash:
            raise HTTPException(status_code=400, detail="Müşteri kodu zaten kullanılıyor")
    for key, value in data.items():
        setattr(customer, key, value)
    customer.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(customer)
    write_audit(
        user_id=user.id,
        action="update",
        entity_type="customer",
        entity_id=customer.id,
        detail={"name": customer.name},
    )
    return _customer_out(customer, _balance_for(db, customer))


@router.delete("/{customer_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_customer(
    customer_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles("admin")),
) -> None:
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Müşteri bulunamadı")
    name = customer.name
    db.delete(customer)
    db.commit()
    write_audit(
        user_id=user.id,
        action="delete",
        entity_type="customer",
        entity_id=customer_id,
        detail={"name": name},
    )


@router.get("/{customer_id}/movements", response_model=list[CariMovementOut])
def list_movements(
    customer_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
    skip: int = 0,
    limit: int = 200,
) -> list[CariMovementOut]:
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Müşteri bulunamadı")
    rows = (
        db.query(CariMovement)
        .filter(CariMovement.customer_id == customer_id)
        .order_by(CariMovement.movement_date.desc(), CariMovement.id.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return [_movement_out(m) for m in rows]


@router.post(
    "/{customer_id}/movements",
    response_model=CariMovementOut,
    status_code=status.HTTP_201_CREATED,
)
def create_movement(
    customer_id: int,
    payload: CariMovementCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*CARI_WRITE_ROLES)),
) -> CariMovementOut:
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Müşteri bulunamadı")
    if payload.movement_type not in CARI_MOVEMENT_TYPES:
        raise HTTPException(status_code=400, detail="Geçersiz hareket tipi")
    if payload.order_id is not None:
        order = db.get(Order, payload.order_id)
        if not order or order.customer_id != customer_id:
            raise HTTPException(status_code=400, detail="Sipariş bu müşteriye ait değil")
    debit, credit = _split_amount(payload.movement_type, payload.amount, payload.side)
    mov_date = payload.movement_date or date.today()
    movement = CariMovement(
        customer_id=customer_id,
        movement_type=payload.movement_type,
        debit=debit,
        credit=credit,
        movement_date=mov_date,
        order_id=payload.order_id,
        note=payload.note,
    )
    db.add(movement)
    db.flush()

    # Nice-to-have: mirror payment/deposit into kasa or banka
    if (
        payload.post_to_finance
        and payload.movement_type in ("payment", "deposit")
        and payload.finance_method
    ):
        from app.models.finance import BankAccount, BankMovement, CashMovement, CashRegister

        note = payload.note or f"Cari {payload.movement_type} — müşteri #{customer_id}"
        if payload.finance_method == "cash":
            reg = (
                db.query(CashRegister)
                .filter(CashRegister.is_active.is_(True))
                .order_by(CashRegister.id.asc())
                .first()
            )
            if reg:
                db.add(
                    CashMovement(
                        cash_register_id=reg.id,
                        movement_type="tahsilat",
                        amount=payload.amount,
                        movement_date=mov_date,
                        note=note,
                        customer_id=customer_id,
                        cari_movement_id=movement.id,
                    )
                )
        elif payload.finance_method == "bank":
            if not payload.bank_account_id:
                raise HTTPException(status_code=400, detail="Banka hesabı seçilmedi")
            acc = db.get(BankAccount, payload.bank_account_id)
            if not acc or not acc.is_active:
                raise HTTPException(status_code=400, detail="Banka hesabı bulunamadı")
            db.add(
                BankMovement(
                    bank_account_id=acc.id,
                    movement_type="deposit",
                    amount=payload.amount,
                    movement_date=mov_date,
                    note=note,
                    customer_id=customer_id,
                    cari_movement_id=movement.id,
                )
            )

    db.commit()
    db.refresh(movement)
    return _movement_out(movement)


@router.get("/{customer_id}/statement", response_model=StatementOut)
def customer_statement(
    customer_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(*READ_ROLES)),
    from_date: date | None = Query(default=None),
    to_date: date | None = Query(default=None),
) -> StatementOut:
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Müşteri bulunamadı")

    query = db.query(CariMovement).filter(CariMovement.customer_id == customer_id)
    if from_date:
        query = query.filter(CariMovement.movement_date >= from_date)
    if to_date:
        query = query.filter(CariMovement.movement_date <= to_date)
    rows = query.order_by(CariMovement.movement_date.asc(), CariMovement.id.asc()).all()

    # Opening for statement period: full opening + movements before from_date
    running = customer.opening_balance or Decimal("0")
    if from_date:
        prior = (
            db.query(
                func.coalesce(func.sum(CariMovement.debit), 0),
                func.coalesce(func.sum(CariMovement.credit), 0),
            )
            .filter(
                CariMovement.customer_id == customer_id,
                CariMovement.movement_date < from_date,
            )
            .one()
        )
        running = running + Decimal(str(prior[0])) - Decimal(str(prior[1]))

    opening = running
    out_rows: list[CariMovementOut] = []
    for m in rows:
        running = running + (m.debit or Decimal("0")) - (m.credit or Decimal("0"))
        out_rows.append(_movement_out(m, running))

    return StatementOut(
        customer_id=customer.id,
        customer_name=customer.name,
        opening_balance=opening,
        closing_balance=running,
        movements=out_rows,
    )


def _finance_method_from_payment_type(payment_type: str) -> tuple[str, str]:
    """Return (finance_method, method_label). finance_method: cash|bank."""
    pt = (payment_type or "Nakit").strip()
    if pt in ("EFT", "Kredi Kartı", "Kredi Karti"):
        return "bank", pt
    return "cash", pt or "Nakit"


def _post_finance_for_tahsilat(
    db: Session,
    *,
    customer_id: int,
    amount: Decimal,
    mov_date: date,
    note: str,
    finance_method: str,
    bank_account_id: int | None,
    cari_movement_id: int | None,
) -> bool:
    from app.models.finance import BankAccount, BankMovement, CashMovement, CashRegister

    if finance_method == "cash":
        reg = (
            db.query(CashRegister)
            .filter(CashRegister.is_active.is_(True))
            .order_by(CashRegister.id.asc())
            .first()
        )
        if not reg:
            return False
        db.add(
            CashMovement(
                cash_register_id=reg.id,
                movement_type="tahsilat",
                amount=amount,
                movement_date=mov_date,
                note=note,
                customer_id=customer_id,
                cari_movement_id=cari_movement_id,
            )
        )
        return True
    if finance_method == "bank":
        acc = None
        if bank_account_id:
            acc = db.get(BankAccount, bank_account_id)
        if acc is None:
            acc = (
                db.query(BankAccount)
                .filter(BankAccount.is_active.is_(True))
                .order_by(BankAccount.id.asc())
                .first()
            )
        if not acc or not acc.is_active:
            raise HTTPException(status_code=400, detail="Aktif banka hesabı bulunamadı")
        db.add(
            BankMovement(
                bank_account_id=acc.id,
                movement_type="deposit",
                amount=amount,
                movement_date=mov_date,
                note=note,
                customer_id=customer_id,
                cari_movement_id=cari_movement_id,
            )
        )
        return True
    return False


def _apply_tahsilat_to_open_orders(
    db: Session,
    customer_id: int,
    amount: Decimal,
    *,
    method: str,
    note: str,
    mov_date: date,
) -> Decimal:
    """Desktop musteri_tahsilatini_acik_siparislere_isle — FIFO open remaining."""
    remaining = amount
    applied = Decimal("0")
    if remaining <= 0:
        return applied
    orders = (
        db.query(Order)
        .options(joinedload(Order.payments))
        .filter(
            Order.customer_id == customer_id,
            ~Order.status.in_(tuple(CLOSED_STATUSES)),
        )
        .order_by(Order.created_at.asc(), Order.id.asc())
        .all()
    )
    for order in orders:
        if remaining <= 0:
            break
        paid = _paid_for_order(order)
        deposit = order.deposit_amount or Decimal("0")
        effective = paid if paid > 0 else deposit
        due = (order.total_amount or Decimal("0")) - effective
        if due <= 0:
            continue
        take = min(remaining, due)
        db.add(
            Payment(
                order_id=order.id,
                amount=take,
                method=method,
                status="tamamlandi",
                paid_at=datetime.combine(mov_date, datetime.min.time()),
                notes=note or f"Müşteri tahsilatı → {order.order_number}",
            )
        )
        remaining -= take
        applied += take
    return applied


@router.post("/{customer_id}/tahsilat", response_model=CustomerTahsilatOut, status_code=status.HTTP_201_CREATED)
def customer_tahsilat(
    customer_id: int,
    payload: CustomerTahsilatIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*CARI_WRITE_ROLES)),
) -> CustomerTahsilatOut:
    """Dedicated Tahsilat Al flow — cari + kasa/banka + açık siparişlere işle."""
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Müşteri bulunamadı")

    lines: list[dict] = [
        {
            "amount": payload.amount,
            "payment_type": payload.payment_type or "Nakit",
            "bank_account_id": payload.bank_account_id,
        }
    ]
    if payload.amount2 and payload.amount2 > 0:
        lines.append(
            {
                "amount": payload.amount2,
                "payment_type": payload.payment_type2 or "Nakit",
                "bank_account_id": payload.bank_account_id2,
            }
        )

    mov_date = payload.movement_date or date.today()
    base_note = (payload.note or "").strip() or "Müşteri tahsilatı"
    movement_ids: list[int] = []
    total = Decimal("0")
    finance_ok = False

    for idx, line in enumerate(lines, start=1):
        amt = Decimal(str(line["amount"]))
        total += amt
        finance_method, method_label = _finance_method_from_payment_type(line["payment_type"])
        note = base_note if len(lines) == 1 else f"{base_note} ({idx}/{len(lines)} · {method_label})"
        movement = CariMovement(
            customer_id=customer_id,
            movement_type="payment",
            debit=Decimal("0"),
            credit=amt,
            movement_date=mov_date,
            note=note,
        )
        db.add(movement)
        db.flush()
        movement_ids.append(movement.id)
        posted = _post_finance_for_tahsilat(
            db,
            customer_id=customer_id,
            amount=amt,
            mov_date=mov_date,
            note=note,
            finance_method=finance_method,
            bank_account_id=line.get("bank_account_id"),
            cari_movement_id=movement.id,
        )
        finance_ok = finance_ok or posted

    applied = Decimal("0")
    if payload.apply_to_open_orders and total > 0:
        primary_method = (payload.payment_type or "Nakit").lower()
        applied = _apply_tahsilat_to_open_orders(
            db,
            customer_id,
            total,
            method=primary_method,
            note=base_note,
            mov_date=mov_date,
        )

    db.commit()
    write_audit(
        user_id=user.id,
        action="create",
        entity_type="customer_tahsilat",
        entity_id=customer_id,
        detail={"amount": float(total), "applied_orders": float(applied)},
    )
    msg = "Tahsilat kaydedildi."
    if applied > 0:
        msg += f" Açık sipariş kalan ödemesine işlenen: {float(applied):,.2f} ₺".replace(",", "X").replace(".", ",").replace("X", ".")
    return CustomerTahsilatOut(
        cari_movement_ids=movement_ids,
        total_amount=total,
        applied_to_orders=applied,
        finance_posted=finance_ok,
        message=msg,
    )


@router.put("/{customer_id}/devir", response_model=CustomerOut)
def customer_devir(
    customer_id: int,
    payload: CustomerDevirIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles(*CARI_WRITE_ROLES)),
) -> CustomerOut:
    """Desktop Devir Bakiye — only opening_balance; does not alter ledger history."""
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(status_code=404, detail="Müşteri bulunamadı")
    amt = abs(Decimal(str(payload.amount)))
    signed = amt if payload.direction == "borclu" else -amt
    customer.opening_balance = signed
    customer.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(customer)
    write_audit(
        user_id=user.id,
        action="update",
        entity_type="customer_devir",
        entity_id=customer.id,
        detail={"opening_balance": float(signed)},
    )
    return _customer_out(customer, _balance_for(db, customer))

